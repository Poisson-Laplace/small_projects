#!/usr/bin/env python3
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

excel_path = './vel_matrix.xlsx'
wb = openpyxl.Workbook()
# def remove
wb.remove(wb.active)

ws37 = wb.create_sheet(title="vel_matrix_37x37")

d_ball = 0.30
x_balls = np.round(np.arange(0.6, 11.41, d_ball), 2) # 37 sütun
z_balls = np.round(np.arange(0.6, 11.41, d_ball), 2) # 37 satır
z_balls_desc = z_balls[::-1]

X37, Z37 = np.meshgrid(x_balls, z_balls_desc)
z_top37 = 8.0 - 0.7 * np.sin(2 * np.pi * X37 / 12.0)
z_bot37 = 4.0 - 0.15 * (X37 - 6.0)

vel37 = np.zeros_like(Z37)
vel37[Z37 >= z_top37] = 1.6
vel37[(Z37 >= z_bot37) & (Z37 < z_top37)] = 2.4
vel37[Z37 < z_bot37] = 3.6
ws37.cell(row=1, column=1, value="z (m) \\ x (m)")

for j, xv in enumerate(x_balls, start=2):
    c = ws37.cell(row=1, column=j, value=float(xv))

for i, zv in enumerate(z_balls_desc, start=2):
    ws37.cell(row=i, column=1, value=float(zv))
    for j, xv in enumerate(x_balls, start=2):
        val = float(vel37[i - 2, j - 2])
        cell = ws37.cell(row=i, column=j, value=val)
        cell.number_format = '0.00'

fill_top = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid") # yellow (1.6 km/s)
fill_mid = PatternFill(start_color="FED7AA", end_color="FED7AA", fill_type="solid") # orange (2.4 km/s)
fill_bot = PatternFill(start_color="93C5FD", end_color="93C5FD", fill_type="solid") # blue (3.6 km/s)
fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
font_header = Font(name="Arial", size=9, bold=True, color="FFFFFF")
font_cell = Font(name="Arial", size=9)

ws37.views.sheetView[0].showGridLines = True
ws37.column_dimensions['A'].width = 12

for j in range(2, len(x_balls) + 2):
    col_letter = get_column_letter(j)
    ws37.column_dimensions[col_letter].width = 6.5

for row in ws37.iter_rows(min_row=1, max_row=len(z_balls) + 1, min_col=1, max_col=len(x_balls) + 1):
    for cell in row:
        if cell.row == 1 or cell.column == 1:
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.font = font_cell
            cell.alignment = Alignment(horizontal="center", vertical="center")
            v = cell.value
            if v is not None:
                if v <= 1.8:
                    cell.fill = fill_top
                elif v <= 2.8:
                    cell.fill = fill_mid
                else:
                    cell.fill = fill_bot

ws_info = wb.create_sheet(title="guide")
ws_info.column_dimensions['A'].width = 25
ws_info.column_dimensions['B'].width = 75


info_rows = [
    ("vel in", "km/s (1.60, 2.40, 3.60 etc.)"),
    ("Format", "2D Matrix (square grid). z up-down."),
    ("Change", "just change the variable."),
    ("Color", "1.6 km/s = yellow, 2.4 km/s = orange, 3.6 km/s = blue."),
    ("Sim", "save the file and done.")
]

ws_info.cell(row=1, column=1, value="feat.").font = font_header
ws_info.cell(row=1, column=1).fill = fill_header
ws_info.cell(row=1, column=2, value="expl.").font = font_header
ws_info.cell(row=1, column=2).fill = fill_header

for r, (k, v) in enumerate(info_rows, start=2):
    ws_info.cell(row=r, column=1, value=k).font = Font(name="Arial", size=10, bold=True)
    ws_info.cell(row=r, column=2, value=v).font = Font(name="Arial", size=10)

wb.save(excel_path)
print(f"Done: {excel_path}")
